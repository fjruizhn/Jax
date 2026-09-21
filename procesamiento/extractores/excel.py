"""Excel → un CSV por hoja.

Una hoja de cálculo NO es texto: aplanarla a prosa destruye justo lo que la
hacía un balance (qué número en qué fila y en qué columna). Por eso el extracto
es CSV por hoja, no un .txt.

openpyxl y no LibreOffice porque LibreOffice convierte una sola hoja
(medido 2026-09-20 sobre un archivo real de 6 hojas: sacó 1, sin error).

`Resultado` (endurecido en a7ef8c9) exige que todo 'ok'/'parcial' tenga AL
MENOS UNA salida con contenido no vacío tras `.strip()`. Un libro cuyas
hojas están TODAS en blanco produciría puros CSV vacíos -- eso violaría esa
regla y `Resultado(...)` levantaría `ValueError`. Por eso ese caso se
detecta ACÁ, antes de construir el Resultado, y se devuelve 'error' con una
razón -- nunca se deja que la excepción se escape de este módulo.

Ronda de arreglo 1 (2026-09-20, task-3-hallazgos.md):
- C-1: la guarda de "todas vacías" NO mira el texto CSV -- una hoja de celdas
  vacías produce ",\\n,\\n" y esas comas sobreviven a `.strip()`, así que la
  guarda vieja nunca disparaba. Se cuenta cuántas CELDAS tienen contenido
  real durante la conversión; si el libro entero suma cero, es 'error'. Una
  hoja sin datos entre otras con datos sigue siendo válida.
- C-2: los `.xlsx` que no vienen de Excel (openpyxl, xlsxwriter, exportadores
  contables) no traen el caché de fórmulas -- con `data_only=True` esa celda
  sale `None`, silencioso. Se abre el libro una SEGUNDA vez con
  `data_only=False`: donde la lectura de valores dio `None` y la cruda es un
  texto que empieza con "=", es una fórmula sin caché. Cuesta leer dos veces;
  perder los números de un balance sin avisar cuesta más.
- I-2: `libro.worksheets` omite las chartsheets -- `hojas` ahora es
  `len(libro.sheetnames)` (lo que el usuario ve), y las no tabulares van a
  `detalle["no_tabulares"]`. Un libro con un gráfico sale 'parcial': algo que
  el usuario ve no se extrajo, y eso hay que decirlo, no callarlo.
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from procesamiento.resultado import Resultado

EXTRACTOR = "openpyxl"


def _version() -> str:
    import openpyxl

    return openpyxl.__version__


def _slug(texto: str) -> str:
    s = re.sub(r"[^\w\-]+", "-", texto.strip(), flags=re.UNICODE).strip("-").lower()
    return s or "hoja"


def _celda_tiene_contenido(valor: object) -> bool:
    if valor is None:
        return False
    if isinstance(valor, str):
        return bool(valor.strip())
    return True


def _hoja_a_csv(hoja) -> tuple[str, int]:
    """CSV de la hoja y cuántas celdas tienen contenido real -- la cuenta
    vive ACÁ, durante la conversión, no se infiere después del texto CSV
    (C-1: una fila de puras celdas vacías igual produce comas)."""
    buffer = io.StringIO()
    escritor = csv.writer(buffer, lineterminator="\n")
    celdas_con_contenido = 0
    for fila in hoja.iter_rows(values_only=True):
        escritor.writerow(["" if c is None else c for c in fila])
        celdas_con_contenido += sum(1 for c in fila if _celda_tiene_contenido(c))
    return buffer.getvalue(), celdas_con_contenido


def _formulas_sin_valor(hoja_valores, hoja_cruda) -> int:
    """Cuenta celdas donde la lectura con `data_only=True` dio `None` pero la
    lectura cruda es un texto de fórmula (empieza con '=') -- una fórmula sin
    caché (C-2)."""
    contador = 0
    for fila_valores, fila_cruda in zip(
        hoja_valores.iter_rows(values_only=True),
        hoja_cruda.iter_rows(values_only=True),
    ):
        for valor, crudo in zip(fila_valores, fila_cruda):
            if valor is None and isinstance(crudo, str) and crudo.startswith("="):
                contador += 1
    return contador


def extraer(origen: Path) -> Resultado:
    import openpyxl

    try:
        libro = openpyxl.load_workbook(origen, data_only=True, read_only=True)
    except Exception as exc:  # archivo corrupto, no-zip, protegido
        return Resultado(
            estado="error", salidas={}, extractor=EXTRACTOR, version=_version(),
            detalle={"razon": f"no se pudo abrir: {type(exc).__name__}: {exc}"},
        )

    try:
        libro_crudo = openpyxl.load_workbook(origen, data_only=False, read_only=True)
    except Exception as exc:
        libro.close()
        return Resultado(
            estado="error", salidas={}, extractor=EXTRACTOR, version=_version(),
            detalle={
                "razon": f"no se pudo abrir para chequeo de formulas: "
                f"{type(exc).__name__}: {exc}",
            },
        )

    nombres_tabulares = [hoja.title for hoja in libro.worksheets]
    set_tabulares = set(nombres_tabulares)
    total = len(libro.sheetnames)
    no_tabulares = [n for n in libro.sheetnames if n not in set_tabulares]
    crudo_por_titulo = {hoja.title: hoja for hoja in libro_crudo.worksheets}

    salidas: dict[str, str] = {}
    fallidas: list[str] = []
    celdas_totales = 0
    formulas_por_hoja: dict[str, int] = {}

    for indice, hoja in enumerate(libro.worksheets, start=1):
        try:
            csv_texto, celdas = _hoja_a_csv(hoja)
        except Exception as exc:
            fallidas.append(f"{hoja.title}: {type(exc).__name__}: {exc}")
            continue

        salidas[f"{indice:02d}-{_slug(hoja.title)}.csv"] = csv_texto
        celdas_totales += celdas

        hoja_cruda = crudo_por_titulo.get(hoja.title)
        if hoja_cruda is not None:
            n_formulas = _formulas_sin_valor(hoja, hoja_cruda)
            if n_formulas:
                formulas_por_hoja[hoja.title] = n_formulas

    libro.close()
    libro_crudo.close()

    if not salidas:
        return Resultado(
            estado="error", salidas={}, extractor=EXTRACTOR, version=_version(),
            detalle={"razon": "ninguna hoja pudo extraerse", "fallidas": fallidas},
        )

    if celdas_totales == 0:
        return Resultado(
            estado="error", salidas={}, extractor=EXTRACTOR, version=_version(),
            detalle={"razon": "ninguna hoja tiene datos", "hojas": total},
        )

    hojas_extraidas = len(salidas)
    total_formulas = sum(formulas_por_hoja.values())
    estado = (
        "ok"
        if hojas_extraidas == total and total_formulas == 0
        else "parcial"
    )

    detalle: dict = {
        "hojas": total,
        "hojas_extraidas": hojas_extraidas,
        "fallidas": fallidas,
    }
    if no_tabulares:
        detalle["no_tabulares"] = no_tabulares
    if total_formulas:
        detalle["formulas_sin_valor"] = {
            "total": total_formulas,
            "hojas": sorted(formulas_por_hoja),
        }

    return Resultado(
        estado=estado, salidas=salidas, extractor=EXTRACTOR, version=_version(),
        detalle=detalle,
    )
