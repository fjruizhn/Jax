"""El defecto que este extractor existe para evitar, reproducido con un archivo
sintético: LibreOffice convierte UNA hoja de seis, calladito.

Ronda de arreglo 1 (2026-09-20, task-3-hallazgos.md): la guarda de "todas las
hojas en blanco" miraba el TEXTO csv (",\\n,\\n" sobrevive a `.strip()`, así que
nunca disparaba). Se corrigió contando celdas con contenido real durante la
conversión (C-1). Además: fórmulas sin valor cacheado se detectan con una
segunda lectura (C-2), las chartsheets ya no desaparecen del conteo de `hojas`
(I-2), y el conjunto de mutaciones se amplió (I-4) porque los tests originales
sólo miraban `hojas`/`hojas_extraidas` en libros donde valen lo mismo.
"""
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

from procesamiento.extractores import excel


def _libro_de_seis_hojas(destino: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i, nombre in enumerate(
        ["2025 JUNIO", "2024", "2023", "2022", "consolidado Bac", "CONSOLIDADO"]
    ):
        ws = wb.create_sheet(title=nombre)
        ws["A1"] = "ACTIVOS"
        ws["B1"] = 100 * (i + 1)
        ws["A2"] = "PASIVOS"
        ws["B2"] = 40 * (i + 1)
    wb.save(destino)
    return destino


def test_un_libro_de_seis_hojas_produce_SEIS_extractos(tmp_path: Path):
    origen = _libro_de_seis_hojas(tmp_path / "eeff.xlsx")
    r = excel.extraer(origen)
    assert r.estado == "ok"
    assert len(r.salidas) == 6, (
        f"se perdieron hojas: {sorted(r.salidas)}. "
        "Este es exactamente el defecto de LibreOffice medido el 2026-09-20."
    )
    assert r.detalle["hojas"] == 6
    assert r.detalle["hojas_extraidas"] == 6


def test_el_contenido_conserva_filas_y_columnas(tmp_path: Path):
    origen = _libro_de_seis_hojas(tmp_path / "eeff.xlsx")
    r = excel.extraer(origen)
    primera = r.salidas[sorted(r.salidas)[0]]
    assert "ACTIVOS,100" in primera.replace("\r\n", "\n")
    assert "PASIVOS,40" in primera.replace("\r\n", "\n")


def test_un_archivo_que_no_es_excel_da_error_sin_extracto(tmp_path: Path):
    malo = tmp_path / "no-es.xlsx"
    malo.write_bytes(b"esto no es un zip")
    r = excel.extraer(malo)
    assert r.estado == "error"
    assert r.salidas == {}
    assert "razon" in r.detalle


def test_un_libro_con_todas_las_celdas_vacias_da_error_incluso_con_filas_definidas(
    tmp_path: Path,
):
    """C-1, el defecto real medido por el auditor: una plantilla donde las
    celdas existen (fueron tocadas -- acá con formato, sin valor) produce
    filas con puras comas (",\\n,\\n"), y ESO sobrevive a `.strip()`. La guarda
    vieja miraba el texto CSV y nunca disparaba. La correcta cuenta celdas con
    contenido real durante la conversión, no el texto de salida."""
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre in ["Hoja A", "Hoja B"]:
        ws = wb.create_sheet(title=nombre)
        for col in ("A", "B"):
            for fila in (1, 2):
                ws[f"{col}{fila}"].font = Font(bold=True)  # celda tocada, SIN valor
    origen = tmp_path / "plantilla-vacia.xlsx"
    wb.save(origen)

    r = excel.extraer(origen)

    assert r.estado == "error"
    assert r.salidas == {}
    assert "razon" in r.detalle
    assert "datos" in r.detalle["razon"]


def test_una_hoja_en_blanco_entre_varias_con_datos_es_un_archivo_valido(tmp_path: Path):
    """El caso legítimo que NO se rechaza: una hoja en blanco entre varias
    con contenido es un archivo válido y tiene que salir 'ok', no 'error'."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws1 = wb.create_sheet(title="con datos")
    ws1["A1"] = "ACTIVOS"
    ws1["B1"] = 100
    wb.create_sheet(title="en blanco")
    ws3 = wb.create_sheet(title="tambien con datos")
    ws3["A1"] = "PASIVOS"
    ws3["B1"] = 40
    origen = tmp_path / "una-hoja-blanco.xlsx"
    wb.save(origen)

    r = excel.extraer(origen)

    assert r.estado == "ok"
    assert len(r.salidas) == 3
    assert r.detalle["hojas"] == 3
    assert r.detalle["hojas_extraidas"] == 3


def test_formula_sin_valor_cacheado_se_detecta_y_pasa_a_parcial(tmp_path: Path):
    """C-2: los .xlsx generados por openpyxl (y varios exportadores contables)
    NO traen el caché de fórmulas. Con `data_only=True` esa celda sale `None`
    -- las etiquetas llegan, los números no, y nada lo decía. Se detecta con
    una segunda lectura cruda y baja el estado a 'parcial'."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "base"
    ws1["A1"] = "ACTIVOS"
    ws1["B1"] = 100

    ws2 = wb.create_sheet("totales")
    ws2["A1"] = "TOTAL ACTIVOS"
    ws2["B1"] = "=base!B1"

    origen = tmp_path / "con-formula.xlsx"
    wb.save(origen)

    r = excel.extraer(origen)

    assert r.estado == "parcial"
    assert r.detalle["formulas_sin_valor"]["total"] == 1
    assert "totales" in r.detalle["formulas_sin_valor"]["hojas"]


def test_excepcion_inesperada_en_formulas_sin_valor_da_resultado_de_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """I-5 (final-hallazgos.md, ronda de cierre): `_formulas_sin_valor()`
    quedaba FUERA de todo `try` -- la lección de D-1 (`word.py`: cualquier
    excepción inesperada del cuerpo sale como `Resultado(estado='error')`,
    nunca cruda) no se había aplicado acá. Sin este arreglo, una excepción
    ahí sube CRUDA hasta `ingerir()` (que no tiene guarda) y de paso fuga
    los dos descriptores del libro, porque `close()` nunca corre."""
    origen = _libro_de_seis_hojas(tmp_path / "eeff.xlsx")

    def _rota(hoja_valores, hoja_cruda):
        raise ValueError("boom en formulas_sin_valor")

    monkeypatch.setattr(excel, "_formulas_sin_valor", _rota)

    r = excel.extraer(origen)

    assert r.estado == "error"
    assert r.salidas == {}
    assert "razon" in r.detalle


def test_excepcion_inesperada_cierra_los_dos_libros_en_finally(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """I-5, la mitad del hallazgo que un `estado='error'` correcto no
    prueba por sí sola: los descriptores de los DOS libros (`libro` y
    `libro_crudo`) tienen que cerrarse aunque el cuerpo reviente -- antes,
    con el `close()` fuera de cualquier `finally`, una excepción los
    dejaba abiertos. Se espía `Workbook.close` para contar cuántas veces
    corrió."""
    origen = _libro_de_seis_hojas(tmp_path / "eeff.xlsx")

    cierres = {"n": 0}
    original_close = openpyxl.workbook.workbook.Workbook.close

    def _close_contado(self):
        cierres["n"] += 1
        return original_close(self)

    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, "close", _close_contado)
    monkeypatch.setattr(
        excel, "_formulas_sin_valor",
        lambda *a, **kw: (_ for _ in ()).throw(ValueError("boom")),
    )

    # Se tolera la excepción acá A PROPÓSITO -- este test aísla el
    # `finally` del `except` que la convierte en `Resultado(estado=
    # "error")` (ese es el otro test, "...da_resultado_de_error"): si sólo
    # se rompe el `finally` (los `close()` sueltos, sin envolver), el
    # `except` puede seguir sano y este test tiene que caer igual.
    try:
        excel.extraer(origen)
    except Exception:  # fail-soft: este test aisla el close() del finally del comportamiento del except que lo convierte en Resultado -- tolera la excepcion cruda a proposito si esa conversion estuviera rota, no es lo que este test verifica
        pass

    assert cierres["n"] == 2, (
        f"I-5 REABIERTO: se esperaban 2 close() (libro + libro_crudo) tras "
        f"una excepcion inesperada, hubo {cierres['n']}"
    )


def test_una_hoja_que_falla_a_extraerse_dejando_las_demas_produce_parcial_con_fallidas(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """I-4: mata MD (`hojas_extraidas` cableado a `total`), M3 (`estado='ok'`
    incondicional) y MI (`fallidas` cableado a `[]`) de una sola vez. NO se
    fuerza corrompiendo el XML -- está medido que openpyxl revienta en
    `load_workbook`, no al iterar, y el libro entero caería a 'error' (razón
    equivocada). Se fuerza con monkeypatch sobre `_hoja_a_csv`."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre in ["buena", "MALA", "otra"]:
        ws = wb.create_sheet(title=nombre)
        ws["A1"] = nombre
        ws["B1"] = 1
    origen = tmp_path / "una-hoja-mala.xlsx"
    wb.save(origen)

    original = excel._hoja_a_csv

    def _rota(hoja):
        if hoja.title == "MALA":
            raise ValueError("boom")
        return original(hoja)

    monkeypatch.setattr(excel, "_hoja_a_csv", _rota)

    r = excel.extraer(origen)

    assert r.estado == "parcial"
    assert sorted(r.salidas) == ["01-buena.csv", "03-otra.csv"]
    assert r.detalle["hojas"] == 3
    assert r.detalle["hojas_extraidas"] == 2
    assert any("MALA" in f for f in r.detalle["fallidas"]), r.detalle["fallidas"]


def test_el_nombre_de_archivo_es_indice_de_dos_digitos_y_slug_exacto(tmp_path: Path):
    """I-4: mata MA (`_slug` siempre 'hoja') y MB (índice sin cero a la
    izquierda) -- el contrato del nombre no estaba aseverado en ningún test."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enero 2025!!"
    ws["A1"] = "x"
    origen = tmp_path / "un-solo-nombre.xlsx"
    wb.save(origen)

    r = excel.extraer(origen)

    assert r.estado == "ok"
    assert set(r.salidas) == {"01-enero-2025.csv"}, sorted(r.salidas)


def test_una_chartsheet_no_desaparece_en_silencio(tmp_path: Path):
    """I-2: `libro.worksheets` omite las hojas de gráfico -- el usuario ve
    N pestañas, nosotros contábamos sólo las tabulares y salíamos 'ok'. Ahora
    `hojas` es lo que el usuario ve (`len(libro.sheetnames)`), y una
    chartsheet fuerza 'parcial' con `no_tabulares` explicando el porqué."""
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "datos"
    ws["A1"] = "x"
    ws["B1"] = 1
    ws["A2"] = "y"
    ws["B2"] = 2

    chart = BarChart()
    data = Reference(ws, min_col=2, min_row=1, max_row=2)
    chart.add_data(data)

    cs = wb.create_chartsheet(title="grafico")
    cs.add_chart(chart)

    origen = tmp_path / "con-grafico.xlsx"
    wb.save(origen)

    r = excel.extraer(origen)

    assert r.estado == "parcial"
    assert r.detalle["hojas"] == 2
    assert r.detalle["hojas_extraidas"] == 1
    assert r.detalle["no_tabulares"] == ["grafico"]


def test_extraer_sin_openpyxl_instalado_da_sin_extractor(tmp_path: Path, monkeypatch):
    """Menor 3 de task-4-hallazgos.md (deuda heredada, pagada junto con el
    mismo defecto en pdf.py): un `ModuleNotFoundError` crudo no es un
    resultado -- existe el estado 'sin_extractor' justo para esto."""
    import sys

    monkeypatch.setitem(sys.modules, "openpyxl", None)
    origen = _libro_de_seis_hojas(tmp_path / "eeff.xlsx")
    r = excel.extraer(origen)
    assert r.estado == "sin_extractor"
    assert r.salidas == {}


def test_version_no_revienta_si_openpyxl_no_esta_instalado(monkeypatch):
    """Ronda P10 (2026-09-21), segundo arreglo del ruling del coordinador:
    `_version()` se llama desde `ingesta._version_vigente` FUERA de
    `extraer()` -- ahí no hay ningún `except ModuleNotFoundError` que
    convierta el fallo en 'sin_extractor'. Antes de este arreglo, `import
    openpyxl` roto acá dejaba escapar un `ImportError` crudo (blindaje que
    `word.py`/`ocr.py` ya tenían y `excel.py` no). Ahora, igual que ellos,
    nunca revienta.

    Menor 10 (final-hallazgos.md, ronda de cierre): el sentinel de "no se
    pudo determinar" es `None`, NUNCA la cadena "desconocida" -- esa cadena
    compara IGUAL A SÍ MISMA en dos fallos consecutivos, y
    `ingesta._version_vigente` la reenvía tal cual para decidir si el
    caché sigue siendo válido (I-2). Contra el código viejo esto falla:
    `_version()` devolvía la cadena "desconocida", no `None`."""
    import sys

    monkeypatch.setitem(sys.modules, "openpyxl", None)
    version = excel._version()
    assert version is None
