"""Word (.docx) → Markdown. Un .docx es un ZIP con XML: el texto, los
títulos y las tablas ya están ahí, estructurados -- no hace falta OCR ni
modelo.

El defecto propio de este extractor: python-docx da la MISMA excepción
(`PackageNotFoundError`, mismo mensaje) para un `.doc` binario viejo (formato
OLE2, que python-docx NUNCA va a poder leer, sin importar cuántas veces se
reintente) que para un `.docx` genuinamente corrupto -- medido a mano
(2026-09-21) antes de escribir el código. Sin distinguir los dos casos, un
`.doc` legítimo que un usuario subió por error de extensión se reportaría
como 'error' (documento roto) en vez de 'sin_extractor' (herramienta
equivocada para este formato) -- la firma OLE2 (`D0 CF 11 E0 A1 B1 1A E1`,
los primeros 8 bytes de CUALQUIER archivo binario de Office viejo: .doc,
.xls, .ppt) se comprueba ANTES de intentar abrir con python-docx.
"""
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")

from procesamiento.extractores import word

_FIRMA_OLE2 = bytes.fromhex("D0CF11E0A1B11AE1")


def _documento(destino: Path) -> Path:
    from docx import Document

    d = Document()
    d.add_heading("Estado de Resultados", level=1)
    d.add_paragraph("Periodo 2026")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "INGRESOS"
    t.cell(0, 1).text = "1000"
    t.cell(1, 0).text = "COSTOS"
    t.cell(1, 1).text = "400"
    d.save(destino)
    return destino


def test_conserva_titulo_parrafo_y_tabla(tmp_path: Path):
    r = word.extraer(_documento(tmp_path / "d.docx"))
    assert r.estado == "ok"
    md = r.salidas["texto.md"]
    assert "# Estado de Resultados" in md
    assert "Periodo 2026" in md
    assert "| INGRESOS | 1000 |" in md


def test_un_docx_roto_da_error_sin_extracto(tmp_path: Path):
    """Bytes al azar -- NO la firma OLE2 -- así que es un archivo
    genuinamente corrupto, no un .doc viejo con la extensión equivocada."""
    malo = tmp_path / "roto.docx"
    malo.write_bytes(b"no soy un docx")
    r = word.extraer(malo)
    assert r.estado == "error"
    assert r.salidas == {}


def test_un_doc_viejo_da_sin_extractor_no_error_criptico(tmp_path: Path):
    """El caso real que python-docx NUNCA va a resolver, y que un mensaje de
    'error' (documento roto) confundiría con un archivo dañado -- acá el
    archivo no está dañado, es el formato equivocado para esta herramienta.
    El contenido tras la firma es basura a propósito: lo único que decide
    la clasificación es la firma OLE2 de los primeros 8 bytes, no el resto."""
    viejo = tmp_path / "balance-2020.doc"
    viejo.write_bytes(_FIRMA_OLE2 + b"\x00" * 200)
    r = word.extraer(viejo)
    assert r.estado == "sin_extractor"
    assert r.salidas == {}
    assert "razon" in r.detalle


def test_un_documento_sin_contenido_da_error_sin_extracto(tmp_path: Path):
    """Fallo cerrado: un .docx válido pero sin un solo párrafo ni tabla con
    texto (el equivalente Word del libro de Excel con todas las hojas en
    blanco, o el PDF sin capa de texto) no puede declararse 'ok' con un
    extracto vacío."""
    from docx import Document

    vacio = tmp_path / "vacio.docx"
    Document().save(vacio)
    r = word.extraer(vacio)
    assert r.estado == "error"
    assert r.salidas == {}


def test_una_celda_con_pipe_no_desalinea_la_tabla(tmp_path: Path):
    """Mismo defecto que pdf.py (I-... del extractor hermano): un '|' dentro
    de una celda se confunde con el separador de columnas de Markdown si no
    se escapa."""
    from docx import Document

    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Ingresos|Egresos"
    t.cell(0, 1).text = "100"
    origen = tmp_path / "con-pipe.docx"
    d.save(origen)

    r = word.extraer(origen)

    assert r.estado == "ok"
    md = r.salidas["texto.md"]
    assert "Ingresos\\|Egresos | 100" in md


def test_una_celda_con_salto_de_linea_no_parte_la_fila(tmp_path: Path):
    """Mismo defecto que pdf.py: una celda de tabla con más de un párrafo
    (un rótulo que envuelve) no puede partir la fila de Markdown en dos."""
    from docx import Document

    d = Document()
    t = d.add_table(rows=1, cols=2)
    celda = t.cell(0, 0)
    celda.text = "Cuentas por"
    celda.add_paragraph("cobrar diversas")
    t.cell(0, 1).text = "450"
    origen = tmp_path / "multilinea.docx"
    d.save(origen)

    r = word.extraer(origen)

    assert r.estado == "ok"
    md = r.salidas["texto.md"]
    assert "Cuentas por cobrar diversas | 450" in md
    # una sola fila de la tabla -> no puede haber una línea de más.
    assert md.count("Cuentas por") == 1


def test_extraer_sin_python_docx_instalado_da_sin_extractor(tmp_path: Path, monkeypatch):
    """Mismo patrón que excel.py y pdf.py: un `ModuleNotFoundError` crudo no
    es un resultado -- existe el estado 'sin_extractor' justo para esto."""
    import sys

    origen = _documento(tmp_path / "d.docx")
    monkeypatch.setitem(sys.modules, "docx", None)
    r = word.extraer(origen)
    assert r.estado == "sin_extractor"
    assert r.salidas == {}
