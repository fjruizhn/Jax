"""Word (.docx) → Markdown. Un .docx es un ZIP con XML: el texto, los
títulos y las tablas ya están ahí, estructurados. No hace falta OCR ni
modelo.

Los .doc viejos (formato binario OLE2) NO los lee python-docx -- nunca,
sin importar la versión ni el reintento. python-docx da la MISMA excepción
(`PackageNotFoundError`, mismo mensaje) para un `.doc` viejo que para un
`.docx` genuinamente corrupto (medido a mano, 2026-09-21) -- así que la
distinción se hace ANTES de intentar abrir el archivo, mirando la firma
OLE2 (`D0 CF 11 E0 A1 B1 1A E1`, los primeros 8 bytes de cualquier binario
de Office viejo: .doc, .xls, .ppt). Con esa firma, es 'sin_extractor'
(herramienta equivocada para este formato, no un documento dañado); sin
ella, si falla al abrir, es 'error' de verdad. Esos van por LibreOffice
(fuera del alcance de la fase 1).
"""
from __future__ import annotations

from pathlib import Path

from procesamiento.resultado import Resultado

EXTRACTOR = "python-docx"

# Firma de las cabeceras OLE2 Compound File Binary Format -- el contenedor
# binario que usan .doc/.xls/.ppt "viejos" (pre-2007). Los primeros 8 bytes
# de CUALQUIER archivo así son EXACTAMENTE estos, sin excepción (es la forma
# en que el formato se identifica a sí mismo). Verificado a mano
# (2026-09-21) contra un .doc real.
_FIRMA_OLE2 = bytes.fromhex("D0CF11E0A1B11AE1")


def _version() -> str:
    try:
        from importlib.metadata import version

        return version("python-docx")
    except Exception:
        return "desconocida"


def _es_binario_ole2(origen: Path) -> bool:
    try:
        with open(origen, "rb") as fh:
            return fh.read(len(_FIRMA_OLE2)) == _FIRMA_OLE2
    except OSError:
        return False


def _tabla_a_lineas(tabla) -> list[str]:
    """Filas de la tabla como bloque Markdown con encabezado -- a diferencia
    de pdf.py (que NUNCA promueve una fila a encabezado, porque una tabla
    detectada por página puede ser la continuación de la anterior), una
    tabla de Word es UN elemento `<w:tbl>` completo en el XML: no hay
    continuación que perder. Un '\\n' dentro de una celda (rótulo que
    envuelve, párrafo extra) se reemplaza por espacio para no partir la
    fila; un '|' se escapa para no desalinear las columnas -- mismos dos
    defectos que ya se pagaron en pdf.py, aplicados acá antes de que
    aparezcan en este extractor."""
    filas = [
        [
            celda.text.strip().replace("\n", " ").replace("|", "\\|")
            for celda in fila.cells
        ]
        for fila in tabla.rows
    ]
    if not filas:
        return []
    lineas = [
        "",
        "| " + " | ".join(filas[0]) + " |",
        "| " + " | ".join(["---"] * len(filas[0])) + " |",
    ]
    for fila in filas[1:]:
        lineas.append("| " + " | ".join(fila) + " |")
    return lineas


def extraer(origen: Path) -> Resultado:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        return Resultado(
            estado="sin_extractor", salidas={}, extractor=EXTRACTOR,
            version="sin instalar",
            detalle={"razon": f"python-docx no esta instalado: {exc}"},
        )

    origen = Path(origen)

    if _es_binario_ole2(origen):
        return Resultado(
            estado="sin_extractor", salidas={}, extractor=EXTRACTOR,
            version=_version(),
            detalle={
                "razon": (
                    "el archivo es un binario OLE2 (.doc/.xls/.ppt de "
                    "formato viejo); python-docx nunca lee ese formato -- "
                    "hace falta LibreOffice, fuera de alcance de esta fase"
                ),
            },
        )

    try:
        documento = Document(str(origen))
    except Exception as exc:
        return Resultado(
            estado="error", salidas={}, extractor=EXTRACTOR, version=_version(),
            detalle={"razon": f"no se pudo abrir: {type(exc).__name__}: {exc}"},
        )

    lineas: list[str] = []
    for parrafo in documento.paragraphs:
        texto = parrafo.text.strip()
        if not texto:
            continue
        estilo = (parrafo.style.name or "").lower()
        if estilo.startswith("heading"):
            nivel = "".join(c for c in estilo if c.isdigit()) or "1"
            lineas.append("#" * min(int(nivel), 6) + " " + texto)
        else:
            lineas.append(texto)

    tablas_con_contenido = 0
    for tabla in documento.tables:
        bloque = _tabla_a_lineas(tabla)
        if bloque:
            tablas_con_contenido += 1
            lineas.extend(bloque)

    contenido = "\n".join(lineas).strip()
    if not contenido:
        return Resultado(
            estado="error", salidas={}, extractor=EXTRACTOR, version=_version(),
            detalle={"razon": "el documento no tiene texto"},
        )

    return Resultado(
        estado="ok", salidas={"texto.md": contenido},
        extractor=EXTRACTOR, version=_version(),
        detalle={
            "parrafos": len(documento.paragraphs),
            "tablas": len(documento.tables),
        },
    )
